"""Multi-format text extraction and data/inputs management.

Reusable by:
  - `core.domain.bank.material` (RAG material indexing).
  - `scripts/index_material.py` (operational CLI).
  - `projects/exam-app` (own attachment upload, stage 2).

Supported formats: PDF (PyMuPDF), DOCX (python-docx), PPTX (python-pptx),
XLSX (openpyxl) and plain text.
"""
from __future__ import annotations

from pathlib import Path


def _extract_pdf(path: str) -> str:
    """Extrae texto de un PDF con PyMuPDF."""
    import fitz  # pymupdf

    doc = fitz.open(path)
    parts = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            parts.append(text.strip())
    doc.close()
    return "\n\n".join(parts)


def _extract_docx(path: str) -> str:
    """Extracts text from a .docx (paragraphs + tables)."""
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(path: str) -> str:
    """Extrae texto de un .pptx (texto de cada shape en cada slide)."""
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_parts = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        slide_parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        slide_parts.append(" | ".join(cells))
        if len(slide_parts) > 1:
            parts.append("\n".join(slide_parts))
    return "\n\n".join(parts)


def _extract_xlsx(path: str) -> str:
    """Extracts text from a .xlsx (all sheets, non-empty cells)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        sheet_parts = [f"--- Hoja: {ws.title} ---"]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                sheet_parts.append(" | ".join(cells))
        if len(sheet_parts) > 1:
            parts.append("\n".join(sheet_parts))
    wb.close()
    return "\n\n".join(parts)


def _extract_text_file(path: str) -> str:
    """Extrae texto de un archivo de texto plano (.txt, .md, etc.)."""
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read().strip()


def extract_file_text(path: str | Path) -> str:
    """Extracts the text of a file according to its extension."""
    path = str(path)
    ext = path.lower().rsplit(".", 1)[-1] if "." in path else ""
    if ext == "pdf":
        return _extract_pdf(path)
    if ext == "docx":
        return _extract_docx(path)
    if ext == "pptx":
        return _extract_pptx(path)
    if ext == "xlsx":
        return _extract_xlsx(path)
    return _extract_text_file(path)
